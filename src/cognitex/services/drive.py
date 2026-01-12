"""Google Drive API integration for document sync."""

import io
from typing import Any, Generator

import structlog
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

from cognitex.services.google_auth import get_google_credentials
from cognitex.config import get_settings

logger = structlog.get_logger()


def get_priority_folders() -> list[str]:
    """Get list of priority folders from config."""
    settings = get_settings()
    folders_str = settings.drive_priority_folders.strip()
    if not folders_str:
        return []
    return [f.strip() for f in folders_str.split(",") if f.strip()]


# Folders to fully index (text extraction + embeddings)
# Now loaded from config, but keep PRIORITY_FOLDERS for backward compatibility
PRIORITY_FOLDERS = get_priority_folders()

# MIME types we can extract text from
EXPORTABLE_MIME_TYPES = {
    "application/vnd.google-apps.document": "text/plain",
    "application/vnd.google-apps.spreadsheet": "text/csv",
    "application/vnd.google-apps.presentation": "text/plain",
}

READABLE_MIME_TYPES = [
    "text/plain",
    "text/markdown",
    "text/csv",
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
]


class DriveService:
    """Service for interacting with Google Drive API."""

    def __init__(self):
        credentials = get_google_credentials()
        self.service = build("drive", "v3", credentials=credentials)

    def list_files(
        self,
        page_size: int = 100,
        query: str | None = None,
        fields: str = "files(id, name, mimeType, modifiedTime, owners, shared, parents, size, webViewLink)",
        order_by: str = "modifiedTime desc",
    ) -> Generator[dict, None, None]:
        """
        List files from Drive with pagination.

        Args:
            page_size: Number of files per page
            query: Optional Drive query string (e.g., "name contains 'report'")
            fields: Fields to return for each file
            order_by: Sort order

        Yields:
            File metadata dictionaries
        """
        page_token = None

        while True:
            results = self.service.files().list(
                pageSize=page_size,
                q=query,
                fields=f"nextPageToken, {fields}",
                orderBy=order_by,
                pageToken=page_token,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
            ).execute()

            for file in results.get("files", []):
                yield file

            page_token = results.get("nextPageToken")
            if not page_token:
                break

    def list_all_files(self, page_size: int = 100) -> Generator[dict, None, None]:
        """List all files (excluding trashed)."""
        yield from self.list_files(
            page_size=page_size,
            query="trashed = false",
        )

    def get_file_metadata(self, file_id: str) -> dict:
        """Get detailed metadata for a specific file."""
        return self.service.files().get(
            fileId=file_id,
            fields="id, name, mimeType, modifiedTime, createdTime, owners, shared, "
                   "sharedWithMeTime, sharingUser, parents, size, webViewLink, "
                   "permissions(emailAddress, role, type)",
            supportsAllDrives=True,
        ).execute()

    def get_folder_id_by_name(self, folder_name: str) -> str | None:
        """
        Find a folder by name (case-insensitive search).

        Returns the folder ID or None if not found.
        """
        query = f"name = '{folder_name}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
        results = self.service.files().list(
            q=query,
            fields="files(id, name)",
            pageSize=10,
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        ).execute()

        files = results.get("files", [])
        if files:
            return files[0]["id"]
        return None

    def list_files_in_folder(
        self,
        folder_id: str,
        recursive: bool = True,
    ) -> Generator[dict, None, None]:
        """
        List all files in a folder, optionally recursively.

        Args:
            folder_id: The folder ID to search
            recursive: If True, include files in subfolders

        Yields:
            File metadata dictionaries with path information
        """
        def _list_recursive(fid: str, path: str = ""):
            # List files in this folder
            query = f"'{fid}' in parents and trashed = false"

            for file in self.list_files(query=query):
                file["_path"] = f"{path}/{file['name']}" if path else file["name"]

                if file["mimeType"] == "application/vnd.google-apps.folder":
                    if recursive:
                        yield from _list_recursive(file["id"], file["_path"])
                else:
                    yield file

        yield from _list_recursive(folder_id)

    def get_file_content(self, file_id: str, mime_type: str) -> str | None:
        """
        Get the text content of a file.

        Args:
            file_id: The file ID
            mime_type: The file's MIME type

        Returns:
            Text content or None if unable to extract
        """
        try:
            # Google Docs/Sheets/Slides - export as text
            if mime_type in EXPORTABLE_MIME_TYPES:
                export_mime = EXPORTABLE_MIME_TYPES[mime_type]
                request = self.service.files().export_media(
                    fileId=file_id,
                    mimeType=export_mime,
                )
                content = request.execute()
                return content.decode("utf-8") if isinstance(content, bytes) else content

            # Regular files - download directly
            if mime_type in READABLE_MIME_TYPES or mime_type.startswith("text/"):
                request = self.service.files().get_media(fileId=file_id)
                buffer = io.BytesIO()
                downloader = MediaIoBaseDownload(buffer, request)

                done = False
                while not done:
                    _, done = downloader.next_chunk()

                content = buffer.getvalue()

                # Handle PDF - would need PyPDF2 or similar
                if mime_type == "application/pdf":
                    return self._extract_pdf_text(content)

                # Handle DOCX - would need python-docx
                if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                    return self._extract_docx_text(content)

                # Handle XLSX - convert to CSV-like text
                if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                    return self._extract_xlsx_text(content)

                # Plain text types
                try:
                    return content.decode("utf-8")
                except UnicodeDecodeError:
                    return content.decode("latin-1", errors="ignore")

            return None

        except Exception as e:
            logger.warning("Failed to get file content", file_id=file_id, error=str(e))
            return None

    def _extract_pdf_text(self, content: bytes) -> str | None:
        """Extract text from PDF bytes."""
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            return "\n\n".join(text_parts) if text_parts else None
        except ImportError:
            logger.warning("pypdf not installed, cannot extract PDF text")
            return None
        except Exception as e:
            logger.warning("Failed to extract PDF text", error=str(e))
            return None

    def _extract_docx_text(self, content: bytes) -> str | None:
        """Extract text from DOCX bytes."""
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            return "\n\n".join([p.text for p in doc.paragraphs if p.text])
        except ImportError:
            logger.warning("python-docx not installed, cannot extract DOCX text")
            return None
        except Exception as e:
            logger.warning("Failed to extract DOCX text", error=str(e))
            return None

    def _extract_xlsx_text(self, content: bytes) -> str | None:
        """Extract text from XLSX bytes as CSV-like format."""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"=== Sheet: {sheet.title} ===")
                for row in sheet.iter_rows(max_row=1000, values_only=True):  # Limit rows
                    row_text = "\t".join(str(c) if c is not None else "" for c in row)
                    if row_text.strip():
                        text_parts.append(row_text)
            wb.close()
            return "\n".join(text_parts) if text_parts else None
        except ImportError:
            logger.warning("openpyxl not installed, cannot extract XLSX text")
            return None
        except Exception as e:
            logger.warning("Failed to extract XLSX text", error=str(e))
            return None

    def get_file_bytes(self, file_id: str, mime_type: str) -> bytes | None:
        """
        Download raw file bytes.

        Args:
            file_id: The file ID
            mime_type: The file's MIME type

        Returns:
            Raw bytes or None if unable to download
        """
        try:
            # Google Docs/Sheets/Slides - export to native format
            if mime_type == "application/vnd.google-apps.document":
                # Export as DOCX for richer analysis
                request = self.service.files().export_media(
                    fileId=file_id,
                    mimeType="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                return request.execute()

            if mime_type == "application/vnd.google-apps.spreadsheet":
                # Export as XLSX
                request = self.service.files().export_media(
                    fileId=file_id,
                    mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                return request.execute()

            if mime_type == "application/vnd.google-apps.presentation":
                # Export as PPTX
                request = self.service.files().export_media(
                    fileId=file_id,
                    mimeType="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
                return request.execute()

            # Regular files - download directly
            request = self.service.files().get_media(fileId=file_id)
            buffer = io.BytesIO()
            downloader = MediaIoBaseDownload(buffer, request)

            done = False
            while not done:
                _, done = downloader.next_chunk()

            return buffer.getvalue()

        except Exception as e:
            logger.warning("Failed to get file bytes", file_id=file_id, error=str(e))
            return None

    async def get_file_analysis(
        self,
        file_id: str,
        file_name: str,
        mime_type: str,
        context: str = "",
    ) -> dict | None:
        """
        Get deep semantic analysis of a file using DocumentAnalyzer.

        Uses Anthropic Skills when available, falls back to local parsing.
        Returns rich analysis including tracked changes, comments, highlights,
        and semantic understanding.

        Args:
            file_id: The file ID
            file_name: The file name
            mime_type: The file's MIME type
            context: Optional context about the file (e.g., "Meeting preparation doc")

        Returns:
            Dict with analysis results or None if unable to analyze
        """
        from cognitex.services.document_analyzer import get_document_analyzer

        # Map Google native types to their export equivalents
        export_mime_map = {
            "application/vnd.google-apps.document": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.google-apps.spreadsheet": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.google-apps.presentation": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
        export_ext_map = {
            "application/vnd.google-apps.document": ".docx",
            "application/vnd.google-apps.spreadsheet": ".xlsx",
            "application/vnd.google-apps.presentation": ".pptx",
        }

        analyzer = get_document_analyzer()

        # Determine effective MIME type and filename for analysis
        effective_mime = export_mime_map.get(mime_type, mime_type)
        if mime_type in export_ext_map and not file_name.endswith(export_ext_map[mime_type]):
            effective_name = file_name + export_ext_map[mime_type]
        else:
            effective_name = file_name

        # Check if document type is supported
        if not analyzer.is_supported(effective_name, effective_mime):
            logger.debug(
                "Document type not supported for deep analysis",
                file_name=file_name,
                mime_type=mime_type,
            )
            return None

        # Download file bytes
        content = self.get_file_bytes(file_id, mime_type)
        if not content:
            logger.warning("Failed to download file for analysis", file_id=file_id)
            return None

        try:
            analysis = await analyzer.analyze(
                filename=effective_name,
                content=content,
                context=context,
                mime_type=effective_mime,
            )

            logger.info(
                "File analyzed successfully",
                file_name=file_name,
                method=analysis.method,
                changes=len(analysis.changes),
                review_items=len(analysis.review_items),
            )

            return analysis.to_dict()

        except Exception as e:
            logger.error("File analysis failed", file_name=file_name, error=str(e))
            return None

    def get_sharing_info(self, file_id: str) -> list[dict]:
        """
        Get detailed sharing information for a file.

        Returns list of sharing entries with email, role, type.
        """
        try:
            permissions = self.service.permissions().list(
                fileId=file_id,
                fields="permissions(emailAddress, role, type, displayName)",
                supportsAllDrives=True,
            ).execute()
            return permissions.get("permissions", [])
        except Exception as e:
            logger.warning("Failed to get sharing info", file_id=file_id, error=str(e))
            return []

    def find_priority_folders(self) -> dict[str, str | None]:
        """
        Find folder IDs for the priority folders.

        Returns dict mapping folder name to folder ID (or None if not found).
        """
        result = {}
        for folder_name in PRIORITY_FOLDERS:
            folder_id = self.get_folder_id_by_name(folder_name)
            result[folder_name] = folder_id
            if folder_id:
                logger.info("Found priority folder", name=folder_name, id=folder_id)
            else:
                logger.warning("Priority folder not found", name=folder_name)
        return result

    def get_start_page_token(self) -> str:
        """Get the starting page token for change tracking."""
        response = self.service.changes().getStartPageToken().execute()
        return response.get('startPageToken')

    def get_changes(self, page_token: str) -> dict:
        """
        Get changes since the given page token.

        Args:
            page_token: Token from previous getStartPageToken or changes.list call

        Returns:
            Dict with 'changes' list and 'newStartPageToken'
        """
        try:
            response = self.service.changes().list(
                pageToken=page_token,
                fields="changes(fileId, removed, file(id, name, mimeType, parents, modifiedTime, trashed)), "
                       "newStartPageToken, nextPageToken",
                includeItemsFromAllDrives=True,
                supportsAllDrives=True,
            ).execute()

            logger.info(
                "Fetched Drive changes",
                change_count=len(response.get('changes', [])),
                has_more=bool(response.get('nextPageToken')),
            )

            return response

        except Exception as e:
            logger.error("Failed to get Drive changes", error=str(e))
            return {'changes': [], 'newStartPageToken': page_token}


# Singleton instance
_drive_service: DriveService | None = None


def get_drive_service() -> DriveService:
    """Get or create the Drive service singleton."""
    global _drive_service
    if _drive_service is None:
        _drive_service = DriveService()
    return _drive_service
